import io
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pdfplumber
import streamlit as st

# 頁面標題與配置
st.set_page_config(
    page_title="牙醫師繼續教育積分整理小工具", layout="centered"
)
st.title("🦷 牙醫師繼續教育積分自動整理器")
st.caption(
    "支援衛福部 PDF 清單：自動篩選「◎參加之專業課程」，並依主辦單位分類統計與匯出"
)

# 檔案上傳元件
uploaded_file = st.file_uploader(
    "請上傳「繼續教育積分及上課紀錄.pdf」", type=["pdf"]
)


def process_cme_pdf(file_bytes):
    full_text = ""
    with pdfplumber.open(file_bytes) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                full_text += t + "\n"

    # 1. 截取「◎參加課程積分」區塊
    start_match = re.search(r"◎\s*參加課程積分", full_text)
    if not start_match:
        return None, "找不到「◎參加課程積分」區塊，請確認上傳的 PDF 格式。"

    sub_text = full_text[start_match.start() :]
    end_matches = list(re.finditer(r"\n\s*◎", sub_text))
    if len(end_matches) > 1:
        sub_text = sub_text[: end_matches[1].start()]

    # 2. 提取專業課程
    lines = [line.strip() for line in sub_text.split("\n") if line.strip()]
    records = []
    current_record = None

    for line in lines:
        if any(
            k in line
            for k in [
                "有效總積分",
                "課程類別",
                "審查單位",
                "衛生福利部",
                "人員類別",
            ]
        ):
            continue

        m1 = re.match(
            r"^專業課程\s+([0-9.]+)\s+([0-9.]+)\s+(\S+)\s+(\S+)\s*(.*)$", line
        )
        m2 = re.match(
            r"^([0-9.]+)\s+([0-9.]+)\s+專業課程\s+(\S+)\s+(\S+)\s*(.*)$", line
        )

        if m1 or m2:
            if current_record:
                records.append(finalize_record(current_record))
            m = m1 if m1 else m2
            v_pts, inv_pts, review, host, rest = m.groups()
            current_record = {
                "有效積分": float(v_pts),
                "無效積分": float(inv_pts),
                "審查單位": review,
                "主辦單位": host,
                "課程名稱": rest,
                "raw_lines": [],
            }
        else:
            if any(
                line.startswith(c)
                for c in ["專業品質", "專業相關法規", "專業倫理"]
            ):
                if current_record:
                    records.append(finalize_record(current_record))
                    current_record = None
                continue
            if current_record:
                current_record["raw_lines"].append(line)

    if current_record:
        records.append(finalize_record(current_record))

    if not records:
        return None, "未解析到任何「專業課程」資料。"

    # 3. 匯出 Excel (使用記憶體 BytesIO)
    output_stream = io.BytesIO()
    build_excel(records, output_stream)
    output_stream.seek(0)
    return output_stream, len(records)


def finalize_record(record):
    raw_tail = " ".join(record.pop("raw_lines"))
    dates = re.findall(r"(\d{4}/\d{1,2}/\d{1,2})", raw_tail)
    course_date = dates[0] if dates else "未知"
    year = course_date.split("/")[0] if dates else "未知"

    cleaned_tail = re.sub(
        r"\d{4}/\d{1,2}/\d{1,2}(?:\s+\d{1,2}:\d{2})?", "", raw_tail
    )
    cleaned_tail = re.sub(r"\b(D1|D2|AG|A|B|C|F|G|H)\b", "", cleaned_tail)
    full_title = re.sub(
        r"\s+", " ", (record["課程名稱"] + " " + cleaned_tail).strip()
    )

    return {
        "主辦單位": record["主辦單位"],
        "課程名稱": full_title,
        "有效積分": record["有效積分"],
        "課程日期": course_date,
        "年度": year,
    }


def build_excel(records, stream):
    org_dict = {}
    for r in records:
        org_dict.setdefault(r["主辦單位"], []).append(r)
    sorted_orgs = sorted(
        org_dict.items(),
        key=lambda x: sum(i["有效積分"] for i in x[1]),
        reverse=True,
    )

    wb = openpyxl.Workbook()
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    double_bottom = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="double", color="1B365D"),
    )

    # 頁籤 1: 總表
    ws1 = wb.active
    ws1.title = "主辦單位統計總表"
    ws1.views.sheetView[0].showGridLines = True
    ws1["A1"] = (
        "牙醫師繼續教育積分 — 參加專業課程統計總表 (依主辦單位分類)"
    )
    ws1["A1"].font = Font(name="微軟正黑體", size=14, bold=True, color="1B365D")
    ws1.append([])
    ws1.append(
        ["主辦單位名稱", "修習堂數", "有效積分加總", "積分佔比 (%)", "課程年份區間"]
    )

    for col in range(1, 6):
        c = ws1.cell(row=3, column=col)
        c.font = Font(name="微軟正黑體", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B365D")
        c.alignment = Alignment(horizontal="center", vertical="center")

    row = 4
    for org, items in sorted_orgs:
        pts = sum(i["有效積分"] for i in items)
        years = sorted(list(set(i["年度"] for i in items)))
        yr_str = f"{years[0]} ~ {years[-1]}" if len(years) > 1 else years[0]

        ws1.cell(row=row, column=1, value=org).font = Font(
            name="微軟正黑體", size=10, bold=True
        )
        ws1.cell(row=row, column=2, value=len(items)).alignment = Alignment(
            horizontal="center"
        )
        ws1.cell(row=row, column=3, value=round(pts, 2)).alignment = Alignment(
            horizontal="right"
        )
        ws1.cell(row=row, column=4, value=f"=C{row}/C{len(sorted_orgs)+4}")
        ws1.cell(row=row, column=5, value=yr_str).alignment = Alignment(
            horizontal="center"
        )
        ws1.cell(row=row, column=3).number_format = "#,##0.00"
        ws1.cell(row=row, column=4).number_format = "0.0%"
        for c in range(1, 6):
            ws1.cell(row=row, column=c).border = thin_border
        row += 1

    # 總計
    ws1.cell(row=row, column=1, value="總計").font = Font(
        name="微軟正黑體", size=10, bold=True, color="1B365D"
    )
    ws1.cell(row=row, column=2, value=f"=SUM(B4:B{row-1})").alignment = (
        Alignment(horizontal="center")
    )
    ws1.cell(row=row, column=3, value=f"=SUM(C4:C{row-1})").alignment = (
        Alignment(horizontal="right")
    )
    ws1.cell(row=row, column=4, value=f"=SUM(D4:D{row-1})").alignment = (
        Alignment(horizontal="right")
    )
    ws1.cell(row=row, column=3).number_format = "#,##0.00"
    ws1.cell(row=row, column=4).number_format = "0.0%"
    for c in range(1, 6):
        ws1.cell(row=row, column=c).border = double_bottom

    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 14
    ws1.column_dimensions["C"].width = 16
    ws1.column_dimensions["D"].width = 16
    ws1.column_dimensions["E"].width = 18

    # 頁籤 2: 明細
    ws2 = wb.create_sheet(title="專業課程明細(依主辦單位分類)")
    ws2.views.sheetView[0].showGridLines = True
    curr_row = 1

    for idx, (org, items) in enumerate(sorted_orgs, 1):
        org_pts = sum(i["有效積分"] for i in items)
        ws2.merge_cells(
            start_row=curr_row, start_column=1, end_row=curr_row, end_column=5
        )
        banner = ws2.cell(
            row=curr_row,
            column=1,
            value=f"【大類 {idx}】 {org} （共 {len(items)} 堂，小計：{org_pts:.2f} 分）",
        )
        banner.font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
        fill_b = (
            PatternFill("solid", fgColor="1B365D")
            if idx % 2 == 1
            else PatternFill("solid", fgColor="006666")
        )
        for c in range(1, 6):
            ws2.cell(row=curr_row, column=c).fill = fill_b
        curr_row += 1

        headers = ["項次", "課程日期", "課程名稱 / 主題", "主辦單位", "有效積分"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws2.cell(row=curr_row, column=c_idx, value=h)
            cell.fill = PatternFill("solid", fgColor="2C4D75")
            cell.font = Font(
                name="微軟正黑體", size=10, bold=True, color="FFFFFF"
            )
            cell.alignment = Alignment(horizontal="center")
        curr_row += 1

        start_r = curr_row
        for item_idx, item in enumerate(items, 1):
            ws2.cell(row=curr_row, column=1, value=item_idx).alignment = (
                Alignment(horizontal="center")
            )
            ws2.cell(row=curr_row, column=2, value=item["課程日期"]).alignment = (
                Alignment(horizontal="center")
            )
            ws2.cell(
                row=curr_row, column=3, value=item["課程名稱"]
            ).alignment = Alignment(wrap_text=True)
            ws2.cell(row=curr_row, column=4, value=org)
            ws2.cell(
                row=curr_row, column=5, value=item["有效積分"]
            ).number_format = "#,##0.00"
            for c in range(1, 6):
                ws2.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1

        # 小計
        ws2.merge_cells(
            start_row=curr_row, start_column=1, end_row=curr_row, end_column=4
        )
        ws2.cell(row=curr_row, column=1, value=f"{org} — 小計").alignment = (
            Alignment(horizontal="right")
        )
        sub_sum = ws2.cell(
            row=curr_row, column=5, value=f"=SUM(E{start_r}:E{curr_row-1})"
        )
        sub_sum.number_format = "#,##0.00"
        sub_sum.font = Font(name="微軟正黑體", bold=True)
        for c in range(1, 6):
            ws2.cell(row=curr_row, column=c).border = thin_border
        curr_row += 2

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 60
    ws2.column_dimensions["D"].width = 24
    ws2.column_dimensions["E"].width = 14

    wb.save(stream)


# 觸發運算與下載
if uploaded_file is not None:
    with st.spinner("正在解析與計算學分..."):
        excel_data, count = process_cme_pdf(uploaded_file)

    if excel_data:
        st.success(f"🎉 處理成功！共擷取到 {count} 堂專業課程。")
        st.download_button(
            label="📥 下載整理好的 Excel 報表",
            data=excel_data,
            file_name="繼續教育專業課程明細_依主辦單位分類.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        st.error(count)